from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from pypdf import PdfReader

EXTRACTOR_VERSION = "extractors-v1"


class ExtractionError(Exception):
    pass


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    extractor: str
    page_count: int | None = None
    warnings: list[str] = field(default_factory=list)
    metadata: dict[str, object] = field(default_factory=dict)


def extract_document(path: Path, mime_type: str) -> ExtractionResult:
    try:
        if mime_type == "application/pdf":
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ExtractionError("Password-protected PDFs are not supported.")
            pages: list[str] = []
            warnings: list[str] = []
            for index, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if not text.strip():
                    warnings.append(f"Page {index + 1} contained no extractable text.")
                pages.append(text)
            return ExtractionResult("\n\n".join(pages), "pypdf", len(reader.pages), warnings)
        if mime_type.endswith("wordprocessingml.document"):
            document = Document(str(path))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
            return ExtractionResult("\n".join(parts), "python-docx")
        if mime_type == "text/plain":
            return ExtractionResult(path.read_text(encoding="utf-8"), "utf-8")
    except ExtractionError:
        raise
    except Exception as error:
        raise ExtractionError("The document could not be parsed.") from error
    raise ExtractionError("Unsupported document type.")
