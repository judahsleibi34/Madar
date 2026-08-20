import math
import ipaddress
import os
import re
import shutil
import socket
import struct
import tempfile
import time
import zipfile
from contextlib import contextmanager
from collections import OrderedDict
from io import BytesIO
from pathlib import Path
from typing import Iterator
from urllib.parse import parse_qs, urljoin, urlparse

import openpyxl
import pandas as pd
import requests

from services.upload_config import (
    get_data_upload_dir,
    resolve_private_user_file_path,
    safe_scope_part,
)


class RemoteDatasetUrlsDisabledError(ValueError):
    pass


class DataReadingNormal:
    SUPPORTED_EXTENSIONS = {".csv", ".xls", ".xlsx"}

    CSV_ENCODINGS = [
        "utf-8-sig",
        "utf-8",
        "cp1256",
        "windows-1256",
        "utf-16",
    ]
    SHARED_CACHE_MAX_ITEMS = int(os.getenv("DATAFRAME_CACHE_MAX_ITEMS", "16"))
    URL_CACHE_SECONDS = int(os.getenv("DATAFRAME_URL_CACHE_SECONDS", "60"))
    MAX_REMOTE_BYTES = int(os.getenv("MAX_REMOTE_DATA_BYTES", str(10 * 1024 * 1024)))
    MAX_CSV_BYTES = int(os.getenv("MAX_CSV_BYTES", os.getenv("MAX_FULL_DATAFRAME_BYTES", str(50 * 1024 * 1024))))
    MAX_ROWS = int(os.getenv("DATAFRAME_MAX_ROWS", "100000"))
    MAX_COLUMNS = int(os.getenv("DATAFRAME_MAX_COLUMNS", "500"))
    MAX_CELL_CHARS = int(os.getenv("DATAFRAME_MAX_CELL_CHARS", "10000"))
    MAX_TOTAL_CELL_CHARS = int(os.getenv("DATAFRAME_MAX_TOTAL_CELL_CHARS", str(10 * 1024 * 1024)))
    MAX_EXCEL_FILE_BYTES = int(os.getenv("MAX_EXCEL_FILE_BYTES", os.getenv("MAX_EXCEL_UPLOAD_BYTES", str(50 * 1024 * 1024))))
    MAX_EXCEL_UNCOMPRESSED_BYTES = int(os.getenv("MAX_EXCEL_UNCOMPRESSED_BYTES", str(50 * 1024 * 1024)))
    MAX_EXCEL_ZIP_ENTRIES = int(os.getenv("MAX_EXCEL_ZIP_ENTRIES", "200"))
    MAX_EXCEL_ZIP_ENTRY_NAME_CHARS = int(os.getenv("MAX_EXCEL_ZIP_ENTRY_NAME_CHARS", "240"))
    MAX_EXCEL_ZIP_COMPRESSION_RATIO = float(os.getenv("MAX_EXCEL_ZIP_COMPRESSION_RATIO", "100"))
    MAX_EXCEL_SHEETS = int(os.getenv("MAX_EXCEL_SHEETS", "20"))
    MAX_EXCEL_ROWS = int(os.getenv("MAX_EXCEL_ROWS", "100000"))
    MAX_EXCEL_COLUMNS = int(os.getenv("MAX_EXCEL_COLUMNS", "1000"))
    MAX_EXCEL_CELL_CHARS = int(os.getenv("MAX_EXCEL_CELL_CHARS", "10000"))
    PARSER_WORKSPACE_BASE = os.getenv("MADAR_UPLOAD_WORKSPACE_DIR", os.path.join(tempfile.gettempdir(), "madar_uploads"))
    REMOTE_DATASET_URLS_DISABLED_MESSAGE = (
        "Remote dataset URLs are disabled. Upload a CSV/XLS/XLSX file instead."
    )
    _shared_df_cache: OrderedDict[str, pd.DataFrame] = OrderedDict()

    def __init__(
        self,
        input_path: str,
        tenant_id: str | int | None = None,
        user_id: str | int | None = None,
    ) -> None:
        if not input_path:
            raise ValueError("input_path is required")

        self.input_path = input_path
        self.tenant_id = tenant_id
        self.user_id = user_id
        self._cached_df: pd.DataFrame | None = None

    def read(self, refresh: bool = False) -> pd.DataFrame:
        if self._is_url(self.input_path):
            self._assert_remote_dataset_urls_enabled()

        if self._cached_df is not None and not refresh:
            return self._cached_df.copy(deep=True)

        try:
            cache_key = None if refresh else self._shared_cache_key()
            if cache_key:
                cached = self._get_shared_cache(cache_key)
                if cached is not None:
                    self._cached_df = cached.copy(deep=True)
                    return cached.copy(deep=True)

            if self._is_google_sheets_url(self.input_path):
                df = self._read_google_sheet(self.input_path)

            elif self._is_url(self.input_path):
                df = self._read_from_url_or_api(self.input_path)

            else:
                df = self._read_local_file(self.input_path)

            self._cached_df = df.copy(deep=True)
            if cache_key:
                self._set_shared_cache(cache_key, df)
            return df.copy(deep=True)

        except Exception as error:
            raise RuntimeError("Failed to read data") from error

    def clear_cache(self) -> None:
        self._cached_df = None

    @classmethod
    def clear_shared_cache(cls) -> None:
        cls._shared_df_cache.clear()

    def _shared_cache_key(self) -> str | None:
        if self._is_url(self.input_path):
            ttl = max(self.URL_CACHE_SECONDS, 1)
            bucket = int(time.time() // ttl)
            return f"url:{bucket}:{self.input_path}"

        safe_path = self._resolve_uploaded_file(self.input_path)
        stat = safe_path.stat()
        return f"file:{safe_path}:{stat.st_size}:{stat.st_mtime_ns}"

    @classmethod
    def _get_shared_cache(cls, key: str) -> pd.DataFrame | None:
        cached = cls._shared_df_cache.get(key)
        if cached is None:
            return None
        cls._shared_df_cache.move_to_end(key)
        return cached.copy(deep=True)

    @classmethod
    def _set_shared_cache(cls, key: str, df: pd.DataFrame) -> None:
        cls._shared_df_cache[key] = df.copy(deep=True)
        cls._shared_df_cache.move_to_end(key)
        while len(cls._shared_df_cache) > cls.SHARED_CACHE_MAX_ITEMS:
            cls._shared_df_cache.popitem(last=False)

    def _read_local_file(self, file_path: str) -> pd.DataFrame:
        safe_path = self._resolve_uploaded_file(file_path)
        extension = self._get_extension(str(safe_path))

        with self._isolated_parser_workspace(safe_path) as workspace_path:
            if extension == ".csv":
                with open(workspace_path, "rb") as file:
                    return self._read_csv_bytes(file.read())

            if extension == ".xlsx":
                self._validate_xlsx_file(workspace_path)
                self._validate_xlsx_workbook(workspace_path)
                df = pd.read_excel(workspace_path, engine="openpyxl")
                self._validate_excel_dataframe(df)
                return self._normalize_dataframe(df)

            if extension == ".xls":
                self._validate_excel_file_size(workspace_path.stat().st_size)
                df = pd.read_excel(workspace_path)
                self._validate_excel_dataframe(df)
                return self._normalize_dataframe(df)

        raise ValueError(
            f"Unsupported file format: {extension}. "
            "Supported formats are .csv, .xls, and .xlsx"
        )

    def _read_from_url_or_api(self, url: str) -> pd.DataFrame:
        extension = self._get_extension(url)

        response = self._fetch_public_url(url)
        response.raise_for_status()

        content_type = response.headers.get("Content-Type", "").lower()

        if extension == ".csv" or "csv" in content_type:
            return self._read_csv_bytes(response.content)

        if extension in [".xls", ".xlsx"] or self._is_excel_content_type(content_type):
            return self._read_excel_bytes(response.content, extension=extension)

        if "application/json" in content_type or url.lower().endswith(".json"):
            data = response.json()
            df = self._json_to_dataframe(data)
            return self._normalize_dataframe(df)

        raise ValueError(
            "Could not detect data format from URL/API. "
            "Expected CSV, Excel, JSON, or Google Sheets."
        )

    def _read_google_sheet(self, url: str) -> pd.DataFrame:
        csv_url = self._google_sheet_to_csv_url(url)

        response = self._fetch_public_url(csv_url)

        if not response.ok:
            raise ValueError(
                "Google Sheet could not be read. Make sure it is shared with "
                "'Anyone with the link' or published to the web."
            )

        content_type = response.headers.get("Content-Type", "").lower()
        content_start = response.content.lstrip()[:20].lower()

        if "html" in content_type or content_start.startswith(b"<"):
            raise ValueError(
                "Google returned a page instead of spreadsheet data. Share the "
                "sheet publicly or publish it to the web, then try again."
            )

        return self._read_csv_bytes(response.content)

    def _read_excel_bytes(self, content: bytes, *, extension: str = "") -> pd.DataFrame:
        self._validate_excel_file_size(len(content))

        if extension == ".xlsx" or zipfile.is_zipfile(BytesIO(content)):
            self._validate_xlsx_file(content)
            self._validate_xlsx_workbook(content)
            df = pd.read_excel(BytesIO(content), engine="openpyxl")
        else:
            df = pd.read_excel(BytesIO(content))

        self._validate_excel_dataframe(df)
        return self._normalize_dataframe(df)

    def _validate_excel_file_size(self, size_bytes: int) -> None:
        if size_bytes > self.MAX_EXCEL_FILE_BYTES:
            raise ValueError("Excel file is too large")

    def _zip_source(self, source):
        if isinstance(source, (bytes, bytearray)):
            return BytesIO(source)
        return source

    def _zip_source_bytes(self, source) -> bytes:
        if isinstance(source, (bytes, bytearray)):
            return bytes(source)

        with open(source, "rb") as file:
            return file.read()

    @contextmanager
    def _isolated_parser_workspace(self, source_path: Path) -> Iterator[Path]:
        # TODO: move parsing into a no-network, resource-limited worker/container.
        # This workspace provides application-level file isolation until then.
        base_dir = Path(self.PARSER_WORKSPACE_BASE).resolve()
        base_dir.mkdir(parents=True, exist_ok=True)
        workspace = Path(
            tempfile.mkdtemp(prefix="parse_", suffix="_job", dir=str(base_dir))
        ).resolve()

        try:
            workspace_file = workspace / f"dataset{source_path.suffix.lower()}"
            shutil.copy2(source_path, workspace_file)
            yield workspace_file
        finally:
            shutil.rmtree(workspace, ignore_errors=True)

    def _validate_xlsx_file(self, source) -> None:
        if not isinstance(source, (bytes, bytearray)):
            self._validate_excel_file_size(Path(source).stat().st_size)
        self._validate_raw_zip_entry_names(source)

        try:
            with zipfile.ZipFile(self._zip_source(source)) as workbook_zip:
                entries = workbook_zip.infolist()

                if len(entries) > self.MAX_EXCEL_ZIP_ENTRIES:
                    raise ValueError("Excel workbook has too many ZIP entries")

                total_uncompressed = 0

                for entry in entries:
                    self._validate_xlsx_zip_entry(entry)
                    total_uncompressed += int(entry.file_size or 0)

                    if total_uncompressed > self.MAX_EXCEL_UNCOMPRESSED_BYTES:
                        raise ValueError("Excel workbook is too large after decompression")

        except zipfile.BadZipFile as error:
            raise ValueError("Invalid or corrupt Excel workbook") from error

    def _validate_xlsx_zip_entry(self, entry: zipfile.ZipInfo) -> None:
        raw_filename = str(getattr(entry, "orig_filename", "") or entry.filename or "")
        filename = str(entry.filename or "")

        if entry.flag_bits & 0x1:
            raise ValueError("Encrypted Excel workbooks are not supported")

        external_file_type = (entry.external_attr >> 16) & 0o170000
        if external_file_type in {0o120000, 0o020000, 0o060000}:
            raise ValueError("Excel workbook contains an unsafe ZIP entry type")

        compressed_size = int(entry.compress_size or 0)
        uncompressed_size = int(entry.file_size or 0)

        if (
            compressed_size > 0
            and uncompressed_size / compressed_size > self.MAX_EXCEL_ZIP_COMPRESSION_RATIO
            and uncompressed_size > 1024 * 1024
        ):
            raise ValueError("Excel workbook has a suspicious compression ratio")

        if uncompressed_size > self.MAX_EXCEL_UNCOMPRESSED_BYTES:
            raise ValueError("Excel workbook is too large after decompression")

        self._validate_xlsx_zip_entry_name(raw_filename)
        self._validate_xlsx_zip_entry_name(filename)

    def _validate_xlsx_zip_entry_name(self, filename: str) -> None:
        if len(filename) > self.MAX_EXCEL_ZIP_ENTRY_NAME_CHARS:
            raise ValueError("Excel workbook contains an unsafe ZIP entry path")

        if (
            not filename
            or filename.startswith(("/", "\\"))
            or "\\" in filename
            or "//" in filename
            or re.match(r"^[A-Za-z]:", filename)
            or ":" in filename.split("/", 1)[0]
            or any(part == "" for part in filename.split("/"))
            or any(part in {".", ".."} for part in filename.split("/"))
        ):
            raise ValueError("Excel workbook contains an unsafe ZIP entry path")

    def _validate_raw_zip_entry_names(self, source) -> None:
        data = self._zip_source_bytes(source)
        for raw_name in self._iter_raw_zip_entry_names(data):
            try:
                filename = raw_name.decode("utf-8")
            except UnicodeDecodeError:
                filename = raw_name.decode("cp437", errors="replace")
            self._validate_xlsx_zip_entry_name(filename)

    def _iter_raw_zip_entry_names(self, data: bytes):
        signatures = (b"PK\x03\x04", b"PK\x01\x02")
        index = 0

        while index < len(data):
            positions = [data.find(signature, index) for signature in signatures]
            positions = [position for position in positions if position >= 0]
            if not positions:
                break

            position = min(positions)
            signature = data[position:position + 4]

            try:
                if signature == b"PK\x03\x04":
                    if position + 30 > len(data):
                        break
                    name_length, extra_length = struct.unpack_from("<HH", data, position + 26)
                    name_start = position + 30
                    next_index = name_start + name_length + extra_length
                elif signature == b"PK\x01\x02":
                    if position + 46 > len(data):
                        break
                    name_length, extra_length, comment_length = struct.unpack_from(
                        "<HHH", data, position + 28
                    )
                    name_start = position + 46
                    next_index = name_start + name_length + extra_length + comment_length
                else:
                    index = position + 1
                    continue

                name_end = name_start + name_length
                if name_length and name_end <= len(data):
                    yield data[name_start:name_end]

                index = max(next_index, position + 4)
            except struct.error:
                break

    def _validate_xlsx_workbook(self, source) -> None:
        try:
            workbook = openpyxl.load_workbook(
                self._zip_source(source),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except zipfile.BadZipFile as error:
            raise ValueError("Invalid or corrupt Excel workbook") from error
        except Exception as error:
            raise ValueError("Invalid or corrupt Excel workbook") from error

        try:
            if len(workbook.sheetnames) > self.MAX_EXCEL_SHEETS:
                raise ValueError("Excel workbook has too many sheets")

            for worksheet in workbook.worksheets:
                if worksheet.max_row and worksheet.max_row > self.MAX_EXCEL_ROWS:
                    raise ValueError("Excel worksheet has too many rows")

                if worksheet.max_column and worksheet.max_column > self.MAX_EXCEL_COLUMNS:
                    raise ValueError("Excel worksheet has too many columns")
        finally:
            workbook.close()

    def _validate_excel_dataframe(self, df: pd.DataFrame) -> None:
        if len(df) > self.MAX_EXCEL_ROWS:
            raise ValueError("Excel worksheet has too many rows")

        if len(df.columns) > self.MAX_EXCEL_COLUMNS:
            raise ValueError("Excel worksheet has too many columns")

        for value in df.to_numpy(dtype=object).flat:
            if isinstance(value, str) and len(value) > self.MAX_EXCEL_CELL_CHARS:
                raise ValueError("Excel cell text is too large")

    def _read_csv_bytes(self, content: bytes) -> pd.DataFrame:
        if not content:
            raise ValueError("CSV file is empty.")

        if len(content) > self.MAX_CSV_BYTES:
            raise ValueError("CSV file is too large.")

        self._reject_binary_csv(content)

        last_error = None

        for encoding in self.CSV_ENCODINGS:
            try:
                df = pd.read_csv(BytesIO(content), encoding=encoding)
                return self._normalize_dataframe(df)
            except UnicodeDecodeError as error:
                last_error = error
            except ValueError:
                raise
            except Exception as error:
                last_error = error

        raise ValueError(f"Could not read CSV with supported encodings: {last_error}")

    def _reject_binary_csv(self, content: bytes) -> None:
        sample = content[:4096]
        if b"\x00" in sample:
            raise ValueError("CSV file appears to be binary.")

        if sample:
            control_bytes = sum(
                1 for byte in sample if byte < 32 and byte not in {9, 10, 13}
            )
            if control_bytes / len(sample) > 0.05:
                raise ValueError("CSV file appears to be binary.")

    def _normalize_dataframe(self, df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()

        df = df.dropna(how="all")
        df = df.dropna(axis=1, how="all")

        if len(df) > self.MAX_ROWS:
            raise ValueError("Dataset has too many rows")

        if len(df.columns) > self.MAX_COLUMNS:
            raise ValueError("Dataset has too many columns")

        total_cell_chars = 0
        for value in df.to_numpy(dtype=object).flat:
            if isinstance(value, str):
                value_length = len(value)
                if value_length > self.MAX_CELL_CHARS:
                    raise ValueError("Dataset cell text is too large")
                total_cell_chars += value_length
                if total_cell_chars > self.MAX_TOTAL_CELL_CHARS:
                    raise ValueError("Dataset text is too large")

        cleaned_columns = []
        seen = {}

        for index, column in enumerate(df.columns):
            name = str(column).replace("\ufeff", "").strip()

            if not name or name.lower().startswith("unnamed:"):
                name = f"Column {index + 1}"

            if name in seen:
                seen[name] += 1
                name = f"{name} {seen[name]}"
            else:
                seen[name] = 1

            cleaned_columns.append(name)

        df.columns = cleaned_columns

        # Replace Pandas missing values with Python None for JSON.
        df = df.astype(object)
        df = df.where(pd.notnull(df), None)

        # Convert problematic values for JSON safety.
        for column in df.columns:
            df[column] = df[column].map(self._json_safe_value)

        return df

    def _json_safe_value(self, value):
        if value is None:
            return None

        if isinstance(value, float):
            if math.isnan(value) or math.isinf(value):
                return None
            return value

        if isinstance(value, pd.Timestamp):
            return value.isoformat()

        return value

    def _google_sheet_to_csv_url(self, url: str) -> str:
        parsed = urlparse(url)

        if parsed.hostname != "docs.google.com":
            raise ValueError("Not a valid Google Sheets URL")

        if "/pubhtml" in parsed.path or parsed.path.endswith("/pub"):
            base_url = url.split("?")[0].replace("/pubhtml", "/pub")
            query = parse_qs(parsed.query)
            gid = query.get("gid", [None])[0]
            return f"{base_url}?output=csv" + (f"&gid={gid}" if gid else "")

        parts = parsed.path.split("/")

        try:
            sheet_id = parts[parts.index("d") + 1]
        except (ValueError, IndexError):
            raise ValueError("Could not extract Google Sheet ID from URL")

        query = parse_qs(parsed.query)
        fragment = parse_qs(parsed.fragment)
        gid = query.get("gid", fragment.get("gid", [None]))[0]

        csv_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv"
        return f"{csv_url}&gid={gid}" if gid else csv_url

    def _json_to_dataframe(self, data) -> pd.DataFrame:
        if isinstance(data, list):
            return pd.DataFrame(data)

        if isinstance(data, dict):
            for value in data.values():
                if isinstance(value, list):
                    return pd.DataFrame(value)

            return pd.DataFrame([data])

        raise ValueError("Unsupported JSON structure")

    def _is_excel_content_type(self, content_type: str) -> bool:
        return (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in content_type
            or "application/vnd.ms-excel" in content_type
        )

    def _get_extension(self, path_or_url: str) -> str:
        path = urlparse(path_or_url).path
        return Path(path).suffix.lower()

    def _is_url(self, value: str) -> bool:
        parsed = urlparse(value)
        return parsed.scheme in ["http", "https"]

    def _is_google_sheets_url(self, value: str) -> bool:
        parsed = urlparse(value)
        return (
            parsed.scheme in ["http", "https"]
            and parsed.hostname == "docs.google.com"
            and "/spreadsheets/" in parsed.path
        )

    def _resolve_uploaded_file(self, file_path: str) -> Path:
        if self.tenant_id is None or self.user_id is None:
            upload_root = get_data_upload_dir().resolve()
            requested = Path(file_path)
            resolved = (Path.cwd() / requested).resolve() if not requested.is_absolute() else requested.resolve()

            if upload_root != resolved and upload_root not in resolved.parents:
                raise ValueError("Only uploaded data files can be read")

            if not resolved.is_file():
                raise ValueError("Uploaded data file was not found")

            return resolved

        try:
            requested = Path(file_path)
            if not requested.is_absolute() and len(requested.parts) == 1:
                file_path = str(
                    self._scoped_upload_root(get_data_upload_dir())
                    / requested.name
                )
            return resolve_private_user_file_path(
                file_path,
                storage_root=get_data_upload_dir(),
                tenant_id=self.tenant_id,
                user_id=self.user_id,
                allowed_extensions=self.SUPPORTED_EXTENSIONS,
            )
        except FileNotFoundError as error:
            raise ValueError("Uploaded data file was not found") from error
        except (PermissionError, ValueError) as error:
            raise ValueError("Only uploaded data files can be read") from error

    def _scoped_upload_root(self, upload_root: Path) -> Path:
        if self.tenant_id is None or self.user_id is None:
            return upload_root

        return (
            upload_root
            / safe_scope_part("tenant", self.tenant_id)
            / safe_scope_part("user", self.user_id)
        ).resolve()

    def _safe_scope_part(self, prefix: str, value: str | int) -> str:
        try:
            return safe_scope_part(prefix, value)
        except ValueError as error:
            raise ValueError("Invalid upload storage scope") from error

    def _fetch_public_url(self, url: str) -> requests.Response:
        self._assert_remote_dataset_urls_enabled()
        headers = {"User-Agent": "Mozilla/5.0"}
        current_url = url

        for _ in range(5):
            self._validate_public_url(current_url)
            response = requests.get(
                current_url,
                timeout=30,
                headers=headers,
                allow_redirects=False,
                stream=True,
            )

            if response.is_redirect:
                response.close()
                location = response.headers.get("Location")
                if not location:
                    raise ValueError("URL redirected without a location")
                current_url = urljoin(current_url, location)
                continue

            self._buffer_limited_response(response)
            return response

        raise ValueError("URL redirected too many times")

    def _buffer_limited_response(self, response: requests.Response) -> None:
        content_length = response.headers.get("Content-Length")

        if content_length:
            try:
                if int(content_length) > self.MAX_REMOTE_BYTES:
                    response.close()
                    raise ValueError("Remote data file is too large")
            except ValueError:
                response.close()
                raise ValueError("Remote data file is too large")

        chunks = []
        total = 0

        for chunk in response.iter_content(chunk_size=1024 * 1024):
            if not chunk:
                continue

            total += len(chunk)

            if total > self.MAX_REMOTE_BYTES:
                response.close()
                raise ValueError("Remote data file is too large")

            chunks.append(chunk)

        response._content = b"".join(chunks)
        response._content_consumed = True

    def _remote_dataset_urls_enabled(self) -> bool:
        return (os.getenv("ALLOW_REMOTE_DATASET_URLS") or "").strip().lower() in {
            "1",
            "true",
            "yes",
            "on",
        }

    def _insecure_remote_dataset_http_enabled(self) -> bool:
        return (os.getenv("ALLOW_INSECURE_REMOTE_DATASET_HTTP") or "").strip().lower() in {
            "1",
            "true",
            "yes",
            "on",
        }

    def _assert_remote_dataset_urls_enabled(self) -> None:
        if not self._remote_dataset_urls_enabled():
            raise RemoteDatasetUrlsDisabledError(self.REMOTE_DATASET_URLS_DISABLED_MESSAGE)

    def _validate_public_url(self, url: str) -> None:
        parsed = urlparse(url)
        if parsed.scheme not in ["http", "https"] or not parsed.hostname:
            raise ValueError("Only public HTTP or HTTPS URLs are supported")

        if parsed.scheme == "http" and not self._insecure_remote_dataset_http_enabled():
            raise ValueError("Remote dataset URLs must use HTTPS")

        try:
            addresses = socket.getaddrinfo(parsed.hostname, parsed.port or 443, type=socket.SOCK_STREAM)
        except socket.gaierror as error:
            raise ValueError("Could not resolve data URL host") from error

        for address in addresses:
            ip = ipaddress.ip_address(address[4][0])
            if (
                ip.is_private
                or ip.is_loopback
                or ip.is_link_local
                or ip.is_multicast
                or ip.is_reserved
                or ip.is_unspecified
            ):
                raise ValueError("Private or local network URLs are not allowed")
