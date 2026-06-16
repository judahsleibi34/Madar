import math
import ipaddress
import os
import re
import socket
import time
import zipfile
from collections import OrderedDict
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlparse

import openpyxl
import pandas as pd
import requests


class RemoteDatasetUrlsDisabledError(ValueError):
    pass


class DataReadingNormal:
    SUPPORTED_EXTENSIONS = {".csv", ".xls", ".xlsx"}

    CSV_ENCODINGS = [
        "utf-8-sig",
        "utf-8",
        "cp1256",
        "windows-1256",
        "utf-16",
    ]
    SHARED_CACHE_MAX_ITEMS = int(os.getenv("DATAFRAME_CACHE_MAX_ITEMS", "16"))
    URL_CACHE_SECONDS = int(os.getenv("DATAFRAME_URL_CACHE_SECONDS", "60"))
    MAX_REMOTE_BYTES = int(os.getenv("MAX_REMOTE_DATA_BYTES", str(10 * 1024 * 1024)))
    MAX_ROWS = int(os.getenv("DATAFRAME_MAX_ROWS", "100000"))
    MAX_COLUMNS = int(os.getenv("DATAFRAME_MAX_COLUMNS", "500"))
    MAX_EXCEL_FILE_BYTES = int(os.getenv("MAX_EXCEL_FILE_BYTES", str(10 * 1024 * 1024)))
    MAX_EXCEL_UNCOMPRESSED_BYTES = int(os.getenv("MAX_EXCEL_UNCOMPRESSED_BYTES", str(50 * 1024 * 1024)))
    MAX_EXCEL_ZIP_ENTRIES = int(os.getenv("MAX_EXCEL_ZIP_ENTRIES", "200"))
    MAX_EXCEL_SHEETS = int(os.getenv("MAX_EXCEL_SHEETS", "20"))
    MAX_EXCEL_ROWS = int(os.getenv("MAX_EXCEL_ROWS", "100000"))
    MAX_EXCEL_COLUMNS = int(os.getenv("MAX_EXCEL_COLUMNS", "1000"))
    MAX_EXCEL_CELL_CHARS = int(os.getenv("MAX_EXCEL_CELL_CHARS", "10000"))
    REMOTE_DATASET_URLS_DISABLED_MESSAGE = (
        "Remote dataset URLs are disabled. Upload a CSV/XLS/XLSX file instead."
    )
    _shared_df_cache: OrderedDict[str, pd.DataFrame] = OrderedDict()

    def __init__(
        self,
        input_path: str,
        tenant_id: str | int | None = None,
        user_id: str | int | None = None,
    ) -> None:
        if not input_path:
            raise ValueError("input_path is required")

        self.input_path = input_path
        self.tenant_id = tenant_id
        self.user_id = user_id
        self._cached_df: pd.DataFrame | None = None

    def read(self, refresh: bool = False) -> pd.DataFrame:
        if self._is_url(self.input_path):
            self._assert_remote_dataset_urls_enabled()

        if self._cached_df is not None and not refresh:
            return self._cached_df.copy(deep=True)

        try:
            cache_key = None if refresh else self._shared_cache_key()
            if cache_key:
                cached = self._get_shared_cache(cache_key)
                if cached is not None:
                    self._cached_df = cached.copy(deep=True)
                    return cached.copy(deep=True)

            if self._is_google_sheets_url(self.input_path):
                df = self._read_google_sheet(self.input_path)

            elif self._is_url(self.input_path):
                df = self._read_from_url_or_api(self.input_path)

            else:
                df = self._read_local_file(self.input_path)

            self._cached_df = df.copy(deep=True)
            if cache_key:
                self._set_shared_cache(cache_key, df)
            return df.copy(deep=True)

        except Exception as error:
            raise RuntimeError("Failed to read data") from error

    def clear_cache(self) -> None:
        self._cached_df = None

    @classmethod
    def clear_shared_cache(cls) -> None:
        cls._shared_df_cache.clear()

    def _shared_cache_key(self) -> str | None:
        if self._is_url(self.input_path):
            ttl = max(self.URL_CACHE_SECONDS, 1)
            bucket = int(time.time() // ttl)
            return f"url:{bucket}:{self.input_path}"

        safe_path = self._resolve_uploaded_file(self.input_path)
        stat = safe_path.stat()
        return f"file:{safe_path}:{stat.st_size}:{stat.st_mtime_ns}"

    @classmethod
    def _get_shared_cache(cls, key: str) -> pd.DataFrame | None:
        cached = cls._shared_df_cache.get(key)
        if cached is None:
            return None
        cls._shared_df_cache.move_to_end(key)
        return cached.copy(deep=True)

    @classmethod
    def _set_shared_cache(cls, key: str, df: pd.DataFrame) -> None:
        cls._shared_df_cache[key] = df.copy(deep=True)
        cls._shared_df_cache.move_to_end(key)
        while len(cls._shared_df_cache) > cls.SHARED_CACHE_MAX_ITEMS:
            cls._shared_df_cache.popitem(last=False)

    def _read_local_file(self, file_path: str) -> pd.DataFrame:
        safe_path = self._resolve_uploaded_file(file_path)
        extension = self._get_extension(str(safe_path))

        if extension == ".csv":
            with open(safe_path, "rb") as file:
                return self._read_csv_bytes(file.read())

        if extension == ".xlsx":
            self._validate_xlsx_file(safe_path)
            self._validate_xlsx_workbook(safe_path)
            df = pd.read_excel(safe_path, engine="openpyxl")
            self._validate_excel_dataframe(df)
            return self._normalize_dataframe(df)

        if extension == ".xls":
            self._validate_excel_file_size(safe_path.stat().st_size)
            df = pd.read_excel(safe_path)
            self._validate_excel_dataframe(df)
            return self._normalize_dataframe(df)

        raise ValueError(
            f"Unsupported file format: {extension}. "
            "Supported formats are .csv, .xls, and .xlsx"
        )

    def _read_from_url_or_api(self, url: str) -> pd.DataFrame:
        extension = self._get_extension(url)

        response = self._fetch_public_url(url)
        response.raise_for_status()

        content_type = response.headers.get("Content-Type", "").lower()

        if extension == ".csv" or "csv" in content_type:
            return self._read_csv_bytes(response.content)

        if extension in [".xls", ".xlsx"] or self._is_excel_content_type(content_type):
            return self._read_excel_bytes(response.content, extension=extension)

        if "application/json" in content_type or url.lower().endswith(".json"):
            data = response.json()
            df = self._json_to_dataframe(data)
            return self._normalize_dataframe(df)

        raise ValueError(
            "Could not detect data format from URL/API. "
            "Expected CSV, Excel, JSON, or Google Sheets."
        )

    def _read_google_sheet(self, url: str) -> pd.DataFrame:
        csv_url = self._google_sheet_to_csv_url(url)

        response = self._fetch_public_url(csv_url)

        if not response.ok:
            raise ValueError(
                "Google Sheet could not be read. Make sure it is shared with "
                "'Anyone with the link' or published to the web."
            )

        content_type = response.headers.get("Content-Type", "").lower()
        content_start = response.content.lstrip()[:20].lower()

        if "html" in content_type or content_start.startswith(b"<"):
            raise ValueError(
                "Google returned a page instead of spreadsheet data. Share the "
                "sheet publicly or publish it to the web, then try again."
            )

        return self._read_csv_bytes(response.content)

    def _read_excel_bytes(self, content: bytes, *, extension: str = "") -> pd.DataFrame:
        self._validate_excel_file_size(len(content))

        if extension == ".xlsx" or zipfile.is_zipfile(BytesIO(content)):
            self._validate_xlsx_file(content)
            self._validate_xlsx_workbook(content)
            df = pd.read_excel(BytesIO(content), engine="openpyxl")
        else:
            df = pd.read_excel(BytesIO(content))

        self._validate_excel_dataframe(df)
        return self._normalize_dataframe(df)

    def _validate_excel_file_size(self, size_bytes: int) -> None:
        if size_bytes > self.MAX_EXCEL_FILE_BYTES:
            raise ValueError("Excel file is too large")

    def _zip_source(self, source):
        if isinstance(source, (bytes, bytearray)):
            return BytesIO(source)
        return source

    def _validate_xlsx_file(self, source) -> None:
        try:
            with zipfile.ZipFile(self._zip_source(source)) as workbook_zip:
                entries = workbook_zip.infolist()

                if len(entries) > self.MAX_EXCEL_ZIP_ENTRIES:
                    raise ValueError("Excel workbook has too many ZIP entries")

                total_uncompressed = 0

                for entry in entries:
                    self._validate_xlsx_zip_entry(entry)
                    total_uncompressed += int(entry.file_size or 0)

                    if total_uncompressed > self.MAX_EXCEL_UNCOMPRESSED_BYTES:
                        raise ValueError("Excel workbook is too large after decompression")

        except zipfile.BadZipFile as error:
            raise ValueError("Invalid or corrupt Excel workbook") from error

    def _validate_xlsx_zip_entry(self, entry: zipfile.ZipInfo) -> None:
        filename = entry.filename or ""

        if entry.flag_bits & 0x1:
            raise ValueError("Encrypted Excel workbooks are not supported")

        if (
            not filename
            or filename.startswith(("/", "\\"))
            or "\\" in filename
            or "//" in filename
            or any(part in {".", ".."} for part in filename.split("/"))
        ):
            raise ValueError("Excel workbook contains an unsafe ZIP entry path")

    def _validate_xlsx_workbook(self, source) -> None:
        try:
            workbook = openpyxl.load_workbook(
                self._zip_source(source),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except zipfile.BadZipFile as error:
            raise ValueError("Invalid or corrupt Excel workbook") from error
        except Exception as error:
            raise ValueError("Invalid or corrupt Excel workbook") from error

        try:
            if len(workbook.sheetnames) > self.MAX_EXCEL_SHEETS:
                raise ValueError("Excel workbook has too many sheets")

            for worksheet in workbook.worksheets:
                if worksheet.max_row and worksheet.max_row > self.MAX_EXCEL_ROWS:
                    raise ValueError("Excel worksheet has too many rows")

                if worksheet.max_column and worksheet.max_column > self.MAX_EXCEL_COLUMNS:
                    raise ValueError("Excel worksheet has too many columns")
        finally:
            workbook.close()

    def _validate_excel_dataframe(self, df: pd.DataFrame) -> None:
        if len(df) > self.MAX_EXCEL_ROWS:
            raise ValueError("Excel worksheet has too many rows")

        if len(df.columns) > self.MAX_EXCEL_COLUMNS:
            raise ValueError("Excel worksheet has too many columns")

        for value in df.to_numpy(dtype=object).flat:
            if isinstance(value, str) and len(value) > self.MAX_EXCEL_CELL_CHARS:
                raise ValueError("Excel cell text is too large")

    def _read_csv_bytes(self, content: bytes) -> pd.DataFrame:
        if not content:
            raise ValueError("CSV file is empty.")

        last_error = None

        for encoding in self.CSV_ENCODINGS:
            try:
                df = pd.read_csv(BytesIO(content), encoding=encoding)
                return self._normalize_dataframe(df)
            except UnicodeDecodeError as error:
                last_error = error
            except Exception as error:
                last_error = error

        raise ValueError(f"Could not read CSV with supported encodings: {last_error}")

    def _normalize_dataframe(self, df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()

        df = df.dropna(how="all")
        df = df.dropna(axis=1, how="all")

        if len(df) > self.MAX_ROWS:
            raise ValueError("Dataset has too many rows")

        if len(df.columns) > self.MAX_COLUMNS:
            raise ValueError("Dataset has too many columns")

        cleaned_columns = []
        seen = {}

        for index, column in enumerate(df.columns):
            name = str(column).replace("\ufeff", "").strip()

            if not name or name.lower().startswith("unnamed:"):
                name = f"Column {index + 1}"

            if name in seen:
                seen[name] += 1
                name = f"{name} {seen[name]}"
            else:
                seen[name] = 1

            cleaned_columns.append(name)

        df.columns = cleaned_columns

        # Replace Pandas missing values with Python None for JSON.
        df = df.astype(object)
        df = df.where(pd.notnull(df), None)

        # Convert problematic values for JSON safety.
        for column in df.columns:
            df[column] = df[column].map(self._json_safe_value)

        return df

    def _json_safe_value(self, value):
        if value is None:
            return None

        if isinstance(value, float):
            if math.isnan(value) or math.isinf(value):
                return None
            return value

        if isinstance(value, pd.Timestamp):
            return value.isoformat()

        return value

    def _google_sheet_to_csv_url(self, url: str) -> str:
        parsed = urlparse(url)

        if parsed.hostname != "docs.google.com":
            raise ValueError("Not a valid Google Sheets URL")

        if "/pubhtml" in parsed.path or parsed.path.endswith("/pub"):
            base_url = url.split("?")[0].replace("/pubhtml", "/pub")
            query = parse_qs(parsed.query)
            gid = query.get("gid", [None])[0]
            return f"{base_url}?output=csv" + (f"&gid={gid}" if gid else "")

        parts = parsed.path.split("/")

        try:
            sheet_id = parts[parts.index("d") + 1]
        except (ValueError, IndexError):
            raise ValueError("Could not extract Google Sheet ID from URL")

        query = parse_qs(parsed.query)
        fragment = parse_qs(parsed.fragment)
        gid = query.get("gid", fragment.get("gid", [None]))[0]

        csv_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv"
        return f"{csv_url}&gid={gid}" if gid else csv_url

    def _json_to_dataframe(self, data) -> pd.DataFrame:
        if isinstance(data, list):
            return pd.DataFrame(data)

        if isinstance(data, dict):
            for value in data.values():
                if isinstance(value, list):
                    return pd.DataFrame(value)

            return pd.DataFrame([data])

        raise ValueError("Unsupported JSON structure")

    def _is_excel_content_type(self, content_type: str) -> bool:
        return (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in content_type
            or "application/vnd.ms-excel" in content_type
        )

    def _get_extension(self, path_or_url: str) -> str:
        path = urlparse(path_or_url).path
        return Path(path).suffix.lower()

    def _is_url(self, value: str) -> bool:
        parsed = urlparse(value)
        return parsed.scheme in ["http", "https"]

    def _is_google_sheets_url(self, value: str) -> bool:
        parsed = urlparse(value)
        return (
            parsed.scheme in ["http", "https"]
            and parsed.hostname == "docs.google.com"
            and "/spreadsheets/" in parsed.path
        )

    def _resolve_uploaded_file(self, file_path: str) -> Path:
        upload_root = Path(os.getenv("DATA_UPLOAD_DIR", "uploads")).resolve()
        requested = Path(file_path)
        resolved = (Path.cwd() / requested).resolve() if not requested.is_absolute() else requested.resolve()
        allowed_root = self._scoped_upload_root(upload_root)

        if allowed_root != resolved and allowed_root not in resolved.parents:
            raise ValueError("Only uploaded data files can be read")

        if not resolved.is_file():
            raise ValueError("Uploaded data file was not found")

        return resolved

    def _scoped_upload_root(self, upload_root: Path) -> Path:
        if self.tenant_id is None or self.user_id is None:
            return upload_root

        return (
            upload_root
            / self._safe_scope_part("tenant", self.tenant_id)
            / self._safe_scope_part("user", self.user_id)
        ).resolve()

    def _safe_scope_part(self, prefix: str, value: str | int) -> str:
        text = str(value).strip()

        if not re.fullmatch(r"[A-Za-z0-9_-]+", text):
            raise ValueError("Invalid upload storage scope")

        return f"{prefix}_{text}"

    def _fetch_public_url(self, url: str) -> requests.Response:
        self._assert_remote_dataset_urls_enabled()
        headers = {"User-Agent": "Mozilla/5.0"}
        current_url = url

        for _ in range(5):
            self._validate_public_url(current_url)
            response = requests.get(
                current_url,
                timeout=30,
                headers=headers,
                allow_redirects=False,
                stream=True,
            )

            if response.is_redirect:
                response.close()
                location = response.headers.get("Location")
                if not location:
                    raise ValueError("URL redirected without a location")
                current_url = urljoin(current_url, location)
                continue

            self._buffer_limited_response(response)
            return response

        raise ValueError("URL redirected too many times")

    def _buffer_limited_response(self, response: requests.Response) -> None:
        content_length = response.headers.get("Content-Length")

        if content_length:
            try:
                if int(content_length) > self.MAX_REMOTE_BYTES:
                    response.close()
                    raise ValueError("Remote data file is too large")
            except ValueError:
                response.close()
                raise ValueError("Remote data file is too large")

        chunks = []
        total = 0

        for chunk in response.iter_content(chunk_size=1024 * 1024):
            if not chunk:
                continue

            total += len(chunk)

            if total > self.MAX_REMOTE_BYTES:
                response.close()
                raise ValueError("Remote data file is too large")

            chunks.append(chunk)

        response._content = b"".join(chunks)
        response._content_consumed = True

    def _remote_dataset_urls_enabled(self) -> bool:
        return (os.getenv("ALLOW_REMOTE_DATASET_URLS") or "").strip().lower() in {
            "1",
            "true",
            "yes",
            "on",
        }

    def _insecure_remote_dataset_http_enabled(self) -> bool:
        return (os.getenv("ALLOW_INSECURE_REMOTE_DATASET_HTTP") or "").strip().lower() in {
            "1",
            "true",
            "yes",
            "on",
        }

    def _assert_remote_dataset_urls_enabled(self) -> None:
        if not self._remote_dataset_urls_enabled():
            raise RemoteDatasetUrlsDisabledError(self.REMOTE_DATASET_URLS_DISABLED_MESSAGE)

    def _validate_public_url(self, url: str) -> None:
        parsed = urlparse(url)
        if parsed.scheme not in ["http", "https"] or not parsed.hostname:
            raise ValueError("Only public HTTP or HTTPS URLs are supported")

        if parsed.scheme == "http" and not self._insecure_remote_dataset_http_enabled():
            raise ValueError("Remote dataset URLs must use HTTPS")

        try:
            addresses = socket.getaddrinfo(parsed.hostname, parsed.port or 443, type=socket.SOCK_STREAM)
        except socket.gaierror as error:
            raise ValueError("Could not resolve data URL host") from error

        for address in addresses:
            ip = ipaddress.ip_address(address[4][0])
            if (
                ip.is_private
                or ip.is_loopback
                or ip.is_link_local
                or ip.is_multicast
                or ip.is_reserved
                or ip.is_unspecified
            ):
                raise ValueError("Private or local network URLs are not allowed")
